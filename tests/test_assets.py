from io import BytesIO

from openpyxl import load_workbook

from app import db
from app.models.asset import Asset, AssetHistory
from app.models.maintenance import Maintenance

from tests.conftest import login


def test_create_asset_and_reject_negative_price(client, admin, catalog):
    login(client)
    rejected = client.post('/asset-management/new', data={
        'tag_number': 'BHSN-T0000099',
        'status': 'Available',
        'asset_type_id': catalog['asset_type_id'],
        'location_id': catalog['location_id'],
        'purchase_price': '-5',
    })
    assert b'cannot be negative' in rejected.data

    created = client.post('/asset-management/new', data={
        'tag_number': 'BHSN-T0000099',
        'serial_number': 'SN-NEW',
        'status': 'Available',
        'asset_type_id': catalog['asset_type_id'],
        'location_id': catalog['location_id'],
        'department_id': catalog['department_id'],
        'manufacturer_id': catalog['manufacturer_id'],
        'purchase_price': '899.50',
    }, follow_redirects=True)
    assert created.status_code == 200
    assert b'BHSN-T0000099' in created.data
    with client.application.app_context():
        asset = Asset.query.filter_by(tag_number='BHSN-T0000099').one()
        assert asset.name == 'Laptop'
        history = AssetHistory.query.filter_by(asset_id=asset.id, action='created').one()
        assert history.changed_by == 'admin'
        assert history.ip_address


def test_duplicate_tag_is_rejected(client, admin, catalog):
    login(client)
    response = client.post('/asset-management/new', data={
        'tag_number': 'BHSN-T0000001',
        'status': 'Available',
        'asset_type_id': catalog['asset_type_id'],
        'location_id': catalog['location_id'],
        'purchase_price': '10',
    })
    assert b'already exists' in response.data


def test_checkout_and_checkin_return_pdf(client, admin, catalog):
    login(client)
    asset_id = catalog['asset_id']
    checkout = client.post(
        f'/asset-management/checkout/{asset_id}',
        data={'employee_id': catalog['employee_id']},
    )
    assert checkout.status_code == 200
    assert checkout.mimetype == 'application/pdf'
    assert checkout.data.startswith(b'%PDF')
    with client.application.app_context():
        asset = db.session.get(Asset, asset_id)
        assert asset.status == 'Checked Out'
        assert asset.current_employee_id == catalog['employee_id']
        assert asset.to_dict()['current_employee'] == 'Jane Doe'

    checkin = client.post(f'/asset-management/checkin/{asset_id}')
    assert checkin.mimetype == 'application/pdf'
    with client.application.app_context():
        asset = db.session.get(Asset, asset_id)
        assert asset.status == 'Available'
        assert asset.current_employee_id is None


def test_audit_updates_timestamp(client, admin, catalog):
    login(client)
    response = client.post(
        f"/asset-management/audit/{catalog['asset_id']}",
        data={
            'condition_status': 'Good',
            'actual_location_id': catalog['location_id'],
            'asset_found': 'yes',
            'serial_verified': 'yes',
            'audit_notes': 'On the shelf',
        },
        follow_redirects=True,
    )
    assert b'No discrepancies found' in response.data
    with client.application.app_context():
        asset = db.session.get(Asset, catalog['asset_id'])
        assert asset.last_audit_date is not None


def test_maintenance_and_dispose(client, admin, catalog):
    login(client)
    asset_id = catalog['asset_id']
    record = client.post(f'/asset-management/maintenance/new/{asset_id}', data={
        'maintenance_type': 'Preventive',
        'priority': 'Low',
        'start_date': '2024-08-01',
        'cost': '40',
        'description': 'Cleaned fans',
    }, follow_redirects=True)
    assert b'Maintenance record created' in record.data
    with client.application.app_context():
        saved = Maintenance.query.filter_by(asset_id=asset_id).one()
        assert float(saved.cost) == 40

    negative = client.post(f'/asset-management/maintenance/new/{asset_id}', data={
        'maintenance_type': 'Corrective',
        'priority': 'High',
        'start_date': '2024-08-02',
        'cost': '-1',
        'description': 'Bad cost',
    })
    assert b'cannot be negative' in negative.data

    sent = client.post(f'/asset-management/maintenance/{asset_id}', follow_redirects=True)
    assert b'sent to maintenance' in sent.data
    blocked = client.post(f'/asset-management/dispose/{asset_id}', follow_redirects=True)
    assert b'currently in maintenance' in blocked.data
    client.post(f'/asset-management/remove-maintenance/{asset_id}')
    disposed = client.post(f'/asset-management/dispose/{asset_id}', follow_redirects=True)
    assert b'disposed successfully' in disposed.data


def test_find_treats_percent_as_literal(client, admin, catalog):
    login(client)
    response = client.get('/asset-management/find?q=%')
    assert response.status_code == 200
    assert b'BHSN-T0000001' not in response.data


def test_filtered_assets_requires_login_and_filters_status(client, admin, catalog):
    anonymous = client.get('/asset-management/api/filtered-assets')
    assert anonymous.status_code == 401
    login(client)
    response = client.get('/asset-management/api/filtered-assets?status=Available')
    payload = response.get_json()
    assert payload['recordsTotal'] == 1
    assert payload['data'][0]['tag_number'] == 'BHSN-T0000001'


def test_export_assets_workbook(client, admin, catalog):
    login(client)
    response = client.get('/asset-management/export')
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.data))
    sheet = workbook.active
    tags = [row[0] for row in sheet.iter_rows(min_row=2, values_only=True)]
    assert 'BHSN-T0000001' in tags


def test_move_redirects_to_edit(client, admin, catalog):
    login(client)
    response = client.get(f"/asset-management/move/{catalog['asset_id']}")
    assert response.status_code == 302
    assert f"/asset-management/{catalog['asset_id']}/edit" in response.headers['Location']


def test_list_search_percent_is_not_match_all(client, admin, catalog):
    login(client)
    response = client.get('/asset-management/list?q=%')
    assert b'BHSN-T0000001' not in response.data
